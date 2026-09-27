import calendar
import json
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_GET, require_http_methods

from masters.models import MstCust, MstCustProd, MstItem, MstItemType, MstProd, MstProdCat
from transactions.constants import QTY_DECIMAL_PLACES

from .lifecycle import build_inventory_lifecycle
from .models import InventoryStock, StockDtl, StockHed, TrnStkAdjHed
from .services import save_stock_adjustment_document
from .stock_adjustment_validation import (
    parse_stock_adjustment_body,
    validate_stock_adjustment_payload,
)


def _parse_flexible_date(raw: str | None, is_end: bool = False) -> date | None:
    """Parse date from multiple possible input formats (ISO, DD-MM-YYYY, DD/MM/YYYY, YYYY-MM, etc.)."""
    if not raw or not str(raw).strip():
        return None
    s = str(raw).strip()
    # Try full date formats
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    # Try month-only formats (e.g. 2026-05, 05-2026, 05/2026)
    for fmt in ('%Y-%m', '%m-%Y', '%m/%Y'):
        try:
            dt = datetime.strptime(s, fmt).date()
            if is_end:
                _, last_day = calendar.monthrange(dt.year, dt.month)
                return date(dt.year, dt.month, last_day)
            return date(dt.year, dt.month, 1)
        except ValueError:
            pass
    return None


def _int_or_none(val) -> int | None:
    if val is None or (isinstance(val, str) and not str(val).strip()):
        return None
    try:
        return int(val)
    except (TypeError, ValueError):
        return None


def _category_id_from_get(raw: str | None) -> str:
    if not raw:
        return ''
    s = str(raw).strip().upper()[:2]
    if len(s) < 1:
        return ''
    if not MstProdCat.objects.filter(pk=s).exists():
        return ''
    return s


from io import BytesIO
from django.http import HttpResponse, JsonResponse
from django.utils.dateparse import parse_date
from .ledger_service import compute_periodic_stock_movements


def export_stock_movement_excel(rows_data: list[dict], period_title: str) -> BytesIO:
    """Build a styled Excel sheet with stock movements."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Stock Movement'

    # Title header
    ws.merge_cells('A1:L1')
    title_cell = ws['A1']
    title_cell.value = f'RENAMED PHARMACEUTICALS - STOCK MOVEMENT REPORT ({period_title})'
    title_cell.font = Font(size=12, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    headers = [
        'Customer', 'Category', 'SKU / Item', 'Type', 'Batch No',
        'Mfg Date', 'Exp Date', 'Opening Qty', 'Receipts (+)', 'Issues (-)', 'Closing Qty', 'UOM'
    ]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 22

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD'),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_idx in (2, 4, 5, 6, 7, 12) else ('right' if col_idx in (8, 9, 10, 11) else 'left'), vertical='center')

    for r_idx, r in enumerate(rows_data, start=4):
        ws.append([
            r['customer'],
            r['category'],
            r['sku'],
            r['sku_kind'],
            r['batch_no'],
            r['mfg_date'],
            r['exp_date'],
            float(r['opening_qty']),
            float(r['receipt_qty']),
            float(r['issue_qty']),
            float(r['closing_qty']),
            r['uom'],
        ])
        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=r_idx, column=c_idx)
            c.border = thin_border
            if c_idx in (8, 9, 10, 11):
                c.number_format = '#,##0.000'
                c.alignment = Alignment(horizontal='right', vertical='center')
            elif c_idx in (2, 4, 5, 6, 7, 12):
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')

    # Auto column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@login_required
@require_GET
def inventory_stock_view(request):
    """Read-only browse of central stock ledger with periodic / monthly movement calculation."""
    customers = MstCust.objects.order_by('cust_name')
    categories = MstProdCat.objects.order_by('prod_cat_name')

    customer_id = _int_or_none(request.GET.get('customer'))
    if customer_id is not None and not MstCust.objects.filter(pk=customer_id).exists():
        customer_id = None

    category_id = _category_id_from_get(request.GET.get('category'))

    product_id = _int_or_none(request.GET.get('product'))
    item_id = _int_or_none(request.GET.get('item'))

    sku_kind = (request.GET.get('sku_kind') or 'all').strip().lower()
    if sku_kind not in ('all', 'product', 'item'):
        sku_kind = 'all'

    batch_q = (request.GET.get('batch') or '').strip()[:100]
    show_zero = request.GET.get('show_zero') == '1'
    export_excel = request.GET.get('export') == 'excel'

    from_date_raw = (request.GET.get('from_date') or '').strip()
    to_date_raw = (request.GET.get('to_date') or '').strip()
    from_date = _parse_flexible_date(from_date_raw, is_end=False)
    to_date = _parse_flexible_date(to_date_raw, is_end=True)

    if from_date and to_date and from_date > to_date:
        from_date, to_date = to_date, from_date

    is_periodic = bool(from_date or to_date)
    use_new_ledger = StockHed.objects.exists()

    if use_new_ledger:
        stock_opts = StockHed.objects.all()
        if customer_id is not None:
            stock_opts = stock_opts.filter(customer_id=customer_id)

        pids = stock_opts.exclude(product__isnull=True).values_list('product_id', flat=True).distinct()
        products = MstProd.objects.filter(pk__in=pids).order_by('prod_name')
        iids = stock_opts.exclude(item__isnull=True).values_list('item_id', flat=True).distinct()
        items = MstItem.objects.filter(pk__in=iids).order_by('item_name')

        allowed_pids = set(products.values_list('pk', flat=True))
        if product_id is not None and product_id not in allowed_pids:
            product_id = None
        allowed_iids = set(items.values_list('pk', flat=True))
        if item_id is not None and item_id not in allowed_iids:
            item_id = None

        qs = StockHed.objects.select_related(
            'customer',
            'product',
            'product__uom',
            'item',
            'item__uom',
            'item_category',
            'financial_year',
        ).order_by('customer__cust_name', 'product__prod_name', 'item__item_name', 'batch_no')

        if customer_id is not None:
            qs = qs.filter(customer_id=customer_id)
        if category_id:
            qs = qs.filter(item_category_id=category_id)
        if product_id is not None:
            qs = qs.filter(product_id=product_id)
        if item_id is not None:
            qs = qs.filter(item_id=item_id)
        if sku_kind == 'product':
            qs = qs.filter(product__isnull=False)
        elif sku_kind == 'item':
            qs = qs.filter(item__isnull=False)
        if batch_q:
            qs = qs.filter(batch_no__icontains=batch_q)

        all_hed_rows = list(qs)
        hed_ids = [r.pk for r in all_hed_rows]

        # Compute movement quantities (either periodic or full ledger running totals)
        if is_periodic:
            period_data = compute_periodic_stock_movements(hed_ids, from_date=from_date, to_date=to_date)
        else:
            period_data = {}

        processed_rows = []
        for r in all_hed_rows:
            if is_periodic:
                mv = period_data.get(r.pk, {})
                op_qty = mv.get('opening_qty', Decimal('0'))
                rc_qty = mv.get('receipt_qty', Decimal('0'))
                is_qty = mv.get('issue_qty', Decimal('0'))
                cl_qty = mv.get('closing_qty', Decimal('0'))
            else:
                op_qty = r.opn_qty
                rc_qty = r.rcpt_qty
                is_qty = r.issue_qty
                cl_qty = r.closing_qty

            r.display_opening_qty = op_qty
            r.display_rcpt_qty = rc_qty
            r.display_issue_qty = is_qty
            r.display_closing_qty = cl_qty

            if not show_zero:
                if cl_qty == Decimal('0') and rc_qty == Decimal('0') and is_qty == Decimal('0') and Decimal(str(r.reserved_qty or 0)) == Decimal('0'):
                    continue

            processed_rows.append(r)

        if export_excel:
            rows_export = []
            for r in processed_rows:
                uom_name = ''
                if r.product_id and r.product.uom:
                    uom_name = r.product.uom.short_name
                elif r.item_id and r.item.uom:
                    uom_name = r.item.uom.short_name
                rows_export.append({
                    'customer': r.customer.cust_name,
                    'category': r.item_category.prod_cat_id if r.item_category else '',
                    'sku': r.product.prod_name if r.product_id else r.item.item_name,
                    'sku_kind': 'FG' if r.product_id else 'RM/PM',
                    'batch_no': r.batch_no or '—',
                    'mfg_date': r.mfg_date.strftime('%d-%m-%Y') if r.mfg_date else '—',
                    'exp_date': r.exp_date.strftime('%d-%m-%Y') if r.exp_date else '—',
                    'opening_qty': r.display_opening_qty,
                    'receipt_qty': r.display_rcpt_qty,
                    'issue_qty': r.display_issue_qty,
                    'closing_qty': r.display_closing_qty,
                    'uom': uom_name,
                })
            period_str = f'{from_date_raw or "Beginning"} to {to_date_raw or "Current"}' if is_periodic else 'All Time / Current'
            buf = export_stock_movement_excel(rows_export, period_str)
            response = HttpResponse(
                buf.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            response['Content-Disposition'] = f'attachment; filename="Stock_Movement_{from_date_raw or "start"}_to_{to_date_raw or "today"}.xlsx"'
            return response

        paginator = Paginator(processed_rows, 50)
        page_num = _int_or_none(request.GET.get('page')) or 1
        page_obj = paginator.get_page(page_num)

        total = len(processed_rows)
        period_label = f' for period {from_date_raw or "start"} to {to_date_raw or "today"}' if is_periodic else ''
        result_summary = f'Showing {len(page_obj.object_list)} row(s) on this page — {total} match(es) in total{period_label}.'
    else:
        stock_opts = InventoryStock.objects.all()
        if customer_id is not None:
            stock_opts = stock_opts.filter(customer_id=customer_id)

        pids = stock_opts.exclude(product__isnull=True).values_list('product_id', flat=True).distinct()
        products = MstProd.objects.filter(pk__in=pids).order_by('prod_name')
        iids = stock_opts.exclude(item__isnull=True).values_list('item_id', flat=True).distinct()
        items = MstItem.objects.filter(pk__in=iids).order_by('item_name')

        allowed_pids = set(products.values_list('pk', flat=True))
        if product_id is not None and product_id not in allowed_pids:
            product_id = None
        allowed_iids = set(items.values_list('pk', flat=True))
        if item_id is not None and item_id not in allowed_iids:
            item_id = None

        qs = InventoryStock.objects.select_related(
            'customer',
            'product',
            'product__uom',
            'item',
            'item__uom',
            'item_category',
            'opened_in_fy',
        ).order_by('customer__cust_name', 'product__prod_name', 'item__item_name', 'batch_no')

        if customer_id is not None:
            qs = qs.filter(customer_id=customer_id)
        if category_id:
            qs = qs.filter(item_category_id=category_id)
        if product_id is not None:
            qs = qs.filter(product_id=product_id)
        if item_id is not None:
            qs = qs.filter(item_id=item_id)
        if sku_kind == 'product':
            qs = qs.filter(product__isnull=False)
        elif sku_kind == 'item':
            qs = qs.filter(item__isnull=False)
        if batch_q:
            qs = qs.filter(batch_no__icontains=batch_q)
        if not show_zero:
            qs = qs.filter(Q(qty__gt=0) | Q(reserved_qty__gt=0))

        all_inv_rows = list(qs)
        for r in all_inv_rows:
            r.display_opening_qty = Decimal('0')
            r.display_rcpt_qty = r.qty
            r.display_issue_qty = Decimal('0')
            r.display_closing_qty = r.qty

        paginator = Paginator(all_inv_rows, 50)
        page_num = _int_or_none(request.GET.get('page')) or 1
        page_obj = paginator.get_page(page_num)

        total = len(all_inv_rows)
        result_summary = f'Showing {len(page_obj.object_list)} row(s) on this page — {total} match(es) in total.'

    q = request.GET.copy()
    q.pop('page', None)
    q.pop('export', None)
    if from_date:
        q['from_date'] = from_date.isoformat()
    if to_date:
        q['to_date'] = to_date.isoformat()
    filter_query = q.urlencode()
    export_query = f'{filter_query}&export=excel' if filter_query else 'export=excel'

    return render(
        request,
        'inventory/inventory_stock_view.html',
        {
            'customers': customers,
            'categories': categories,
            'products': products,
            'items': items,
            'page_obj': page_obj,
            'filter_query': filter_query,
            'export_query': export_query,
            'result_summary': result_summary,
            'sel_customer_id': customer_id,
            'sel_category_id': category_id,
            'sel_product_id': product_id,
            'sel_item_id': item_id,
            'sel_sku_kind': sku_kind,
            'sel_batch': batch_q,
            'sel_show_zero': show_zero,
            'sel_from_date': from_date.isoformat() if from_date else '',
            'sel_to_date': to_date.isoformat() if to_date else '',
            'sel_from_date_display': from_date.strftime('%d-%m-%Y') if from_date else '',
            'sel_to_date_display': to_date.strftime('%d-%m-%Y') if to_date else '',
            'is_periodic': is_periodic,
            'using_stock_hed': use_new_ledger,
            'total': total,
        },
    )


@login_required
@require_GET
def inventory_stock_lifecycle_ajax(request):
    """Batch-wise movement history for one inventory ledger row (JSON)."""
    inv_id = _int_or_none(request.GET.get('inv_id'))
    if inv_id is None:
        return JsonResponse({'error': 'inv_id is required.'}, status=400)

    # Check StockHed first
    hed_row = StockHed.objects.select_related('customer', 'product', 'item').filter(pk=inv_id).first()
    if hed_row:
        movements = list(hed_row.movements.all().order_by('trn_date', 'stock_dtl_id'))
        events = []
        running_bal = Decimal('0')
        step = Decimal('1').scaleb(-QTY_DECIMAL_PLACES)
        for m in movements:
            if m.trn_type in (StockDtl.TYPE_OPENING, StockDtl.TYPE_RECEIPT):
                delta = -m.quantity if m.is_reversal else m.quantity
            else:
                delta = m.quantity if m.is_reversal else -m.quantity
            running_bal = (running_bal + delta).quantize(step, rounding=ROUND_HALF_UP)
            rev_tag = ' (Reversal)' if m.is_reversal else ''
            events.append({
                'trn_date': m.trn_date.isoformat(),
                'type_key': m.tran_id,
                'label': f'{m.get_tran_id_display()}{rev_tag}',
                'qty_delta': str(delta),
                'qty_delta_display': f'{delta:+f}',
                'balance_after': str(running_bal),
                'doc_type': m.tran_id,
                'doc_id': m.trn_no,
                'doc_display': f'{m.tran_id} #{m.trn_no}',
                'detail': m.remarks or '',
                'doc_url': '',
            })

        sku_name = hed_row.product.prod_name if hed_row.product else hed_row.item.item_name
        uom_str = ''
        if hed_row.product and hed_row.product.uom:
            uom_str = hed_row.product.uom.short_name
        elif hed_row.item and hed_row.item.uom:
            uom_str = hed_row.item.uom.short_name

        payload = {
            'inv_id': hed_row.stock_lnkno,
            'customer_name': hed_row.customer.cust_name,
            'sku_name': sku_name,
            'sku_kind': 'product' if hed_row.product_id else 'item',
            'batch_no': hed_row.batch_no or '—',
            'uom': uom_str,
            'ledger_qty': str(hed_row.closing_qty),
            'computed_balance': str(running_bal),
            'balance_matches_ledger': hed_row.closing_qty == running_bal,
            'events': events,
            'item_level_events': [],
            'ledger_episodes': [],
            'note': '',
        }
        return JsonResponse(payload)

    row = get_object_or_404(
        InventoryStock.objects.select_related('customer', 'product', 'item'),
        pk=inv_id,
    )
    try:
        payload = build_inventory_lifecycle(
            customer_id=row.customer_id,
            product_id=row.product_id,
            item_id=row.item_id,
            batch_no=row.batch_no,
        )
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    payload['inv_id'] = row.inv_id
    return JsonResponse(payload)


def _stock_adj_document_dict(hed: TrnStkAdjHed) -> dict:
    lines = []
    for ln in hed.lines.prefetch_related('batches').select_related('product', 'item'):
        kind = 'P' if ln.product_id else 'I'
        mid = ln.product_id or ln.item_id
        label = ln.product.prod_name if kind == 'P' else ln.item.item_name
        batches = []
        for b in ln.batches.all():
            batches.append(
                {
                    'batch_no': b.batch_no,
                    'mfg_date': b.mfg_date.isoformat() if b.mfg_date else '',
                    'exp_date': b.exp_date.isoformat() if b.exp_date else '',
                    'batch_qty': str(b.batch_qty),
                }
            )
        lines.append(
            {
                'kind': kind,
                'master_id': mid,
                'label': label,
                'remarks': ln.line_remarks,
                'batches': batches,
            }
        )
    return {
        'stk_adj_id': hed.stk_adj_id,
        'customer_id': hed.customer_id,
        'stk_adj_dt': hed.stk_adj_dt.isoformat(),
        'stk_adj_type': hed.stk_adj_type,
        'item_type_id': hed.item_type_id,
        'remarks': hed.remarks or '',
        'lines': lines,
    }


@login_required
@require_GET
def stock_adjustment_masters_ajax(request):
    """Return product or item rows for the selected customer + item type."""
    try:
        customer_id = int(request.GET.get('customer_id') or 0)
    except (TypeError, ValueError):
        customer_id = 0
    try:
        item_type_id = int(request.GET.get('item_type_id') or 0)
    except (TypeError, ValueError):
        item_type_id = 0
    if customer_id <= 0 or item_type_id <= 0:
        return JsonResponse({'mode': '', 'rows': []})

    cust_prods = (
        MstCustProd.objects.filter(customer_id=customer_id, product__prod_type_id=item_type_id)
        .select_related('product')
        .order_by('product__prod_name')
    )
    if cust_prods.exists():
        rows = [{'id': cp.product_id, 'label': cp.product.prod_name} for cp in cust_prods]
        return JsonResponse({'mode': 'P', 'rows': rows})

    items = MstItem.objects.filter(item_type_id=item_type_id).order_by('item_name')
    rows = [{'id': it.item_id, 'label': it.item_name} for it in items]
    return JsonResponse({'mode': 'I', 'rows': rows})


@login_required
@require_GET
def stock_adjustment_sku_meta_ajax(request):
    """
    Mfg/exp field enablement from item master (Y/N flags). Products (FG) always allow dates.
    """
    try:
        customer_id = int(request.GET.get('customer_id') or 0)
    except (TypeError, ValueError):
        customer_id = 0
    try:
        master_id = int(request.GET.get('master_id') or 0)
    except (TypeError, ValueError):
        master_id = 0
    kind = (request.GET.get('kind') or '').strip().upper()
    if customer_id <= 0 or master_id <= 0 or kind not in ('P', 'I'):
        return JsonResponse({'maintain_batch': 'N', 'mfg_enabled': True, 'exp_enabled': True})
    if kind == 'P':
        if not MstCustProd.objects.filter(customer_id=customer_id, product_id=master_id).exists():
            return JsonResponse({'error': 'not_found'}, status=404)
        return JsonResponse({'maintain_batch': 'Y', 'mfg_enabled': True, 'exp_enabled': True})
    it = MstItem.objects.filter(pk=master_id).first()
    if not it:
        return JsonResponse({'error': 'not_found'}, status=404)
    return JsonResponse(
        {
            'maintain_batch': it.maintain_batch or 'N',
            'mfg_enabled': it.mfg_date == 'Y',
            'exp_enabled': it.exp_date == 'Y',
        }
    )


@login_required
@require_GET
def stock_adjustment_inv_defaults_ajax(request):
    """Suggest mfg/exp from existing ledger row for this customer + SKU + batch (optional)."""
    try:
        customer_id = int(request.GET.get('customer_id') or 0)
    except (TypeError, ValueError):
        customer_id = 0
    try:
        master_id = int(request.GET.get('master_id') or 0)
    except (TypeError, ValueError):
        master_id = 0
    kind = (request.GET.get('kind') or '').strip().upper()
    batch_no = (request.GET.get('batch_no') or '').strip()[:100]
    if customer_id <= 0 or master_id <= 0 or kind not in ('P', 'I'):
        return JsonResponse({'mfg_date': '', 'exp_date': ''})

    flt_hed: dict = {'customer_id': customer_id}
    if kind == 'P':
        flt_hed['product_id'] = master_id
        flt_hed['batch_no'] = batch_no
        flt_hed['item_id__isnull'] = True
    else:
        flt_hed['item_id'] = master_id
        flt_hed['batch_no'] = ''
        flt_hed['product_id__isnull'] = True

    row = StockHed.objects.filter(**flt_hed, is_closed=False).first()
    if row and (row.mfg_date or row.exp_date):
        return JsonResponse(
            {
                'mfg_date': row.mfg_date.isoformat() if row.mfg_date else '',
                'exp_date': row.exp_date.isoformat() if row.exp_date else '',
            }
        )

    flt_old: dict = {'customer_id': customer_id, 'batch_no': batch_no}
    if kind == 'P':
        flt_old['product_id'] = master_id
        flt_old['item_id'] = None
    else:
        flt_old['item_id'] = master_id
        flt_old['product_id'] = None
    old_row = InventoryStock.objects.filter(**flt_old, is_closed=False).first()
    if not old_row:
        return JsonResponse({'mfg_date': '', 'exp_date': ''})
    return JsonResponse(
        {
            'mfg_date': old_row.mfg_date.isoformat() if old_row.mfg_date else '',
            'exp_date': old_row.exp_date.isoformat() if old_row.exp_date else '',
        }
    )


@login_required
@require_http_methods(['GET', 'POST'])
def stock_adjustment_view(request):
    customers = MstCust.objects.order_by('cust_name')
    item_types = MstItemType.objects.select_related('item_category').order_by('item_type_name')

    try:
        edit_pk = int(request.GET.get('edit_pk') or request.POST.get('edit_pk') or 0)
    except (TypeError, ValueError):
        edit_pk = 0

    initial_doc = None
    if edit_pk:
        hed = get_object_or_404(TrnStkAdjHed, pk=edit_pk)
        initial_doc = _stock_adj_document_dict(hed)

    repost_doc = None

    if request.method == 'POST':
        raw = request.POST.get('stk_adj_document') or ''
        e1, body = parse_stock_adjustment_body(raw)
        if e1:
            for msg in e1:
                messages.error(request, msg)
            repost_doc = raw or None
        else:
            try:
                edit_pk_post = int(request.POST.get('edit_pk') or 0)
            except (TypeError, ValueError):
                edit_pk_post = 0
            e2, cleaned = validate_stock_adjustment_payload(body, edit_pk=edit_pk_post or None)
            if e2:
                for msg in e2[:25]:
                    messages.error(request, msg)
                if len(e2) > 25:
                    messages.error(request, 'Fix the errors above (additional messages omitted).')
                repost_doc = raw
            else:
                assert cleaned is not None
                try:
                    hed = save_stock_adjustment_document(
                        header_pk=cleaned['header_pk'],
                        customer_id=cleaned['customer_id'],
                        stk_adj_dt=cleaned['stk_adj_dt'],
                        stk_adj_type=cleaned['stk_adj_type'],
                        item_type_id=cleaned['item_type_id'],
                        remarks=cleaned['remarks'],
                        lines_payload=cleaned['lines'],
                    )
                except ValidationError as exc:
                    for msg in exc.messages:
                        messages.error(request, msg)
                    repost_doc = raw
                else:
                    messages.success(request, 'Stock adjustment saved.')
                    return redirect('inventory:stock_adjustment')

    display_initial = initial_doc
    if repost_doc:
        try:
            display_initial = json.loads(repost_doc)
        except json.JSONDecodeError:
            pass

    return render(
        request,
        'inventory/stock_adjustment.html',
        {
            'customers': customers,
            'item_types': item_types,
            'initial_doc': display_initial,
            'edit_pk': edit_pk or None,
            'today': date.today().isoformat(),
        },
    )
