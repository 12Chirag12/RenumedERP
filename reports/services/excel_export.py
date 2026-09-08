"""Excel export for tabular master report payloads."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def _display_width(value) -> int:
    """Approximate max line length for column sizing (handles newlines)."""
    if value is None:
        return 0
    text = str(value)
    if not text:
        return 0
    return max(len(line) for line in text.splitlines())


def build_workbook(*, title: str, column_defs: list[dict], rows: list[dict]) -> BytesIO:
    """
    column_defs: [{'key': 'x', 'label': 'X'}, ...]
    rows: list of dicts with those keys (values may be None).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = (title or 'Report')[:31]

    headers = [c['label'] for c in column_defs]
    keys = [c['key'] for c in column_defs]
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical='top')

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical='top')

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, key in enumerate(keys, start=1):
            val = row.get(key)
            if val is None:
                val = ''
            c = ws.cell(row=r_idx, column=c_idx, value=val)
            c.alignment = wrap

    # Column widths from content (openpyxl width ≈ character count for default font)
    for c_idx, key in enumerate(keys, start=1):
        letter = get_column_letter(c_idx)
        max_len = max(_display_width(headers[c_idx - 1]), 8)
        for r_idx in range(2, ws.max_row + 1):
            v = ws.cell(row=r_idx, column=c_idx).value
            max_len = max(max_len, _display_width(v))
        # Cap very wide columns; wrap_text prevents overlap within the cell
        ws.column_dimensions[letter].width = min(max(max_len + 2.5, 10), 85)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_product_attributes_workbook(*, title: str, sections: list) -> BytesIO:
    """
    One sheet: each attribute section in its own column pair (side by side),
    with a merged title row (e.g. Product category, Coating type).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = (title or 'Product attributes')[:31]

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical='top')
    center_title = Alignment(wrap_text=True, vertical='center', horizontal='center')

    gap_cols = 1
    start_col = 1

    for sec in sections or []:
        end_col = start_col + 1
        sec_title = sec.get('title') or sec.get('table_key') or ''
        col_defs = sec.get('column_defs') or [
            {'key': 'sr', 'label': 'Sr. No.'},
            {'key': 'name', 'label': 'Value'},
        ]
        h_left = col_defs[0].get('label') or 'Sr. No.'
        h_right = col_defs[1].get('label') or 'Value'

        ws.merge_cells(
            start_row=1, start_column=start_col, end_row=1, end_column=end_col
        )
        tcell = ws.cell(row=1, column=start_col, value=sec_title)
        tcell.font = bold
        tcell.alignment = center_title

        r = 2
        if sec.get('layout') == 'grouped':
            for g in sec.get('groups') or []:
                ws.merge_cells(
                    start_row=r, start_column=start_col, end_row=r, end_column=end_col
                )
                gh = ws.cell(row=r, column=start_col, value=g.get('header') or '')
                gh.font = bold
                gh.alignment = wrap
                r += 1
                for ci, lab in enumerate((h_left, h_right), start=start_col):
                    c = ws.cell(row=r, column=ci, value=lab)
                    c.font = bold
                    c.alignment = wrap
                r += 1
                for row in g.get('rows') or []:
                    ws.cell(row=r, column=start_col, value=row.get('sr')).alignment = wrap
                    ws.cell(
                        row=r, column=start_col + 1, value=row.get('name')
                    ).alignment = wrap
                    r += 1
        else:
            for ci, lab in enumerate((h_left, h_right), start=start_col):
                c = ws.cell(row=r, column=ci, value=lab)
                c.font = bold
                c.alignment = wrap
            r += 1
            for row in sec.get('rows') or []:
                ws.cell(row=r, column=start_col, value=row.get('sr')).alignment = wrap
                ws.cell(
                    row=r, column=start_col + 1, value=row.get('name')
                ).alignment = wrap
                r += 1

        for c_idx in range(start_col, end_col + 1):
            letter = get_column_letter(c_idx)
            max_len = max(_display_width(sec_title), _display_width(h_left), _display_width(h_right), 10)
            for rr in range(1, ws.max_row + 1):
                v = ws.cell(row=rr, column=c_idx).value
                max_len = max(max_len, _display_width(v))
            ws.column_dimensions[letter].width = min(max(max_len + 1.5, 12), 48)

        start_col = end_col + 1 + gap_cols

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
